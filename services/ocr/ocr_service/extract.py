"""Turn an uploaded file into plain text.

Supported: PDF (text layer, OCR for pages without one), images (PNG, JPEG,
TIFF incl. multi-page, WebP, GIF, BMP), DOCX, XLSX, and plain-text formats
(text, Markdown, CSV, JSON, XML, YAML, source code) and HTML. Anything else,
including RTF, is refused (HTTP 415), never passed on unread.

The extracted text is capped (``OCR_MAX_TEXT_CHARS``) while it is built, so
a small file that expands to hundreds of megabytes of text is refused.
"""

from __future__ import annotations

import io
import math
import zipfile
from dataclasses import dataclass

from PIL import Image, ImageOps

from .config import Settings
from .engines import OcrEngine
from .text import fix_hungarian_accents, html_to_text, tidy

IMAGE_FORMATS = ["PNG", "JPEG", "TIFF", "WEBP", "GIF", "BMP"]
# Plain-text formats are passed through as they are, so the list is closed:
# a format with its own markup or escapes (RTF writes "P\'e9ter") would hide
# names from the recognizers.
TEXT_TYPES = {
    "text/plain", "text/markdown", "text/x-markdown", "text/csv", "text/tab-separated-values",
    "application/json", "text/xml", "application/xml", "text/yaml", "application/yaml",
    "application/x-yaml", "text/x-python", "text/javascript", "application/javascript", "text/x-log",
}
TEXT_EXTENSIONS = {
    "txt", "md", "markdown", "csv", "tsv", "json", "log", "xml", "yaml", "yml", "ini", "cfg", "toml",
    "py", "js", "ts", "java", "c", "h", "cpp", "cs", "go", "rs", "sql", "sh",
}
HTML_TYPES = {"text/html", "application/xhtml+xml"}
HTML_EXTENSIONS = {"html", "htm", "xhtml"}
# Cells read from a spreadsheet; an empty row counts as one, so a sheet with
# an absurd row index cannot loop for hours.
MAX_SPREADSHEET_CELLS = 500_000
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class ExtractionError(Exception):
    """The file could not be read (damaged, encrypted, wrong type): HTTP 422."""


class UnsupportedTypeError(Exception):
    """File type not supported: HTTP 415."""


class TooLargeError(Exception):
    """Over a page, pixel, size or text limit: HTTP 413."""


class EngineFailure(Exception):
    """The OCR engine itself failed (e.g. its model server is down): HTTP 503."""


@dataclass
class Extracted:
    text: str
    kind: str
    pages: int = 0
    ocr_pages: int = 0


class TextBudget:
    """Counts extracted characters and stops as soon as the limit is passed."""

    def __init__(self, limit: int) -> None:
        self.limit = limit
        self.used = 0

    def add(self, text: str) -> str:
        self.used += len(text)
        if self.used > self.limit:
            raise TooLargeError(f"the file contains more than {self.limit} characters of text")
        return text


def sniff(data: bytes, content_type: str, filename: str) -> str:
    """Decide the file kind from its first bytes; the declared type and the
    file name extension are only used for text formats and BMP."""
    ctype = content_type.split(";", 1)[0].strip().lower()
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    head = data[:16]
    if head.startswith(b"%PDF-"):
        return "pdf"
    if (
        head.startswith(b"\x89PNG\r\n\x1a\n")
        or head.startswith(b"\xff\xd8\xff")
        or head[:4] in (b"II*\x00", b"MM\x00*")
        or (head[:4] == b"RIFF" and head[8:12] == b"WEBP")
        or head[:6] in (b"GIF87a", b"GIF89a")
        or (head[:2] == b"BM" and (ctype == "image/bmp" or ext == "bmp"))
    ):
        return "image"
    if head.startswith(b"PK\x03\x04"):
        try:
            names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
        except zipfile.BadZipFile as exc:
            raise ExtractionError("damaged ZIP container") from exc
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        raise UnsupportedTypeError("ZIP-based format not supported (only DOCX and XLSX)")
    if data.lstrip()[:5] == b"{\\rtf":
        raise UnsupportedTypeError("RTF is not supported; save the document as DOCX or PDF")
    if ctype in HTML_TYPES or ext in HTML_EXTENSIONS:
        return "html"
    if ctype in TEXT_TYPES or ext in TEXT_EXTENSIONS:
        return "text"
    raise UnsupportedTypeError(f"file type not supported: {ctype or 'unknown'}")


def extract(data: bytes, content_type: str, filename: str, engine: OcrEngine, settings: Settings) -> Extracted:
    kind = sniff(data, content_type, filename)
    budget = TextBudget(settings.max_text_chars)
    try:
        if kind == "pdf":
            result = _pdf(data, engine, settings, budget)
        elif kind == "image":
            result = _image(data, engine, settings, budget)
        elif kind == "docx":
            _check_zip(data, settings)
            result = Extracted(_docx(data, budget), kind)
        elif kind == "xlsx":
            _check_zip(data, settings)
            result = Extracted(_xlsx(data, budget), kind)
        elif kind == "html":
            result = Extracted(html_to_text(budget.add(_decode(data))), kind)
        else:
            result = Extracted(budget.add(_decode(data)), kind)
    except (ExtractionError, TooLargeError, UnsupportedTypeError, EngineFailure):
        raise
    except Image.DecompressionBombError as exc:
        raise TooLargeError("image exceeds the pixel limit") from exc
    except Exception as exc:  # noqa: BLE001 - any other parser failure means the file is unreadable
        raise ExtractionError(f"the file cannot be read ({type(exc).__name__})") from exc
    result.text = tidy(result.text)
    return result


def _decode(data: bytes) -> str:
    if b"\x00" in data:
        raise UnsupportedTypeError("binary data declared as text")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1250", errors="replace")  # legacy Hungarian Windows encoding


def _ocr(image: Image.Image, engine: OcrEngine, settings: Settings, dpi: int, budget: TextBudget) -> str:
    try:
        text = engine.ocr([image], dpi)[0]
    except Exception as exc:  # noqa: BLE001
        raise EngineFailure(type(exc).__name__) from exc
    return budget.add(fix_hungarian_accents(text) if settings.fix_hu_accents else text)


# -- PDF ------------------------------------------------------------------------
def _pdf(data: bytes, engine: OcrEngine, settings: Settings, budget: TextBudget) -> Extracted:
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as exc:
        raise ExtractionError("PDF cannot be opened (damaged or password-protected)") from exc
    try:
        n = len(pdf)
        if n > settings.max_pages:
            raise TooLargeError(f"PDF has {n} pages; the limit is {settings.max_pages}")
        texts: list[str] = []
        ocr_pages = 0
        for i in range(n):
            page = pdf[i]
            try:
                textpage = page.get_textpage()
                layer = textpage.get_text_bounded()
                textpage.close()
                if sum(not c.isspace() for c in layer) >= settings.min_text_layer_chars:
                    texts.append(budget.add(layer))
                    continue
                width, height = page.get_size()  # points (1/72 inch)
                scale = settings.pdf_dpi / 72
                if width * height * scale * scale > settings.max_pixels:
                    scale = math.sqrt(settings.max_pixels / (width * height))
                bitmap = page.render(scale=scale)
                image = bitmap.to_pil()
                bitmap.close()
            finally:
                page.close()
            texts.append(_ocr(image, engine, settings, round(scale * 72), budget))
            ocr_pages += 1
        return Extracted("\n\n".join(texts), "pdf", pages=n, ocr_pages=ocr_pages)
    finally:
        pdf.close()


# -- images -----------------------------------------------------------------------
def _image(data: bytes, engine: OcrEngine, settings: Settings, budget: TextBudget) -> Extracted:
    try:
        image = Image.open(io.BytesIO(data), formats=IMAGE_FORMATS)
    except (OSError, Image.DecompressionBombError) as exc:
        raise ExtractionError("image cannot be decoded") from exc
    if image.width * image.height > settings.max_pixels:
        raise TooLargeError(f"image has {image.width}x{image.height} pixels; the limit is {settings.max_pixels}")
    dpi = _dpi(image.info.get("dpi"))
    # Count frames by seeking at most max_pages+1 times: Pillow's n_frames
    # walks every frame of a TIFF, quadratically.
    frames = 0
    for i in range(settings.max_pages + 1):
        try:
            image.seek(i)
        except EOFError:
            break
        frames = i + 1
    if frames > settings.max_pages:
        raise TooLargeError(f"image has more than {settings.max_pages} frames")
    texts: list[str] = []
    # One frame at a time: a multi-page TIFF or GIF is never held in memory
    # as a whole.
    for i in range(frames):
        image.seek(i)
        if image.width * image.height > settings.max_pixels:
            raise TooLargeError(f"image frame exceeds {settings.max_pixels} pixels")
        frame = ImageOps.exif_transpose(image.convert("RGB"))
        texts.append(_ocr(frame, engine, settings, dpi, budget))
        del frame
    return Extracted("\n\n".join(texts), "image", pages=frames, ocr_pages=frames)


def _dpi(value) -> int:
    """Resolution from image metadata, clamped; 300 when absent or malformed."""
    try:
        dpi = float(value[0] if isinstance(value, (tuple, list)) else value)
    except (TypeError, ValueError, IndexError, ZeroDivisionError):
        return 300
    return int(min(max(dpi, 70), 1200)) if dpi == dpi and dpi > 0 else 300


# -- Office -------------------------------------------------------------------------
def _check_zip(data: bytes, settings: Settings) -> None:
    """Refuse ZIP bombs before a parser inflates them."""
    try:
        infos = zipfile.ZipFile(io.BytesIO(data)).infolist()
    except zipfile.BadZipFile as exc:
        raise ExtractionError("damaged ZIP container") from exc
    if sum(i.file_size for i in infos) > settings.max_unzipped_bytes:
        raise TooLargeError("document expands beyond the size limit")


def _docx(data: bytes, budget: TextBudget) -> str:
    """Visible text of a Word document in reading order: every header (first
    page, even, default), the body, every footer, then footnotes and
    endnotes. Includes tables, content controls, tracked insertions and text
    boxes. Leaves out tracked deletions and comments, which are not part of
    the visible document."""
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import parse_xml

    document = docx.Document(io.BytesIO(data))
    parts_by_type: dict[str, list] = {RT.HEADER: [], RT.FOOTER: [], RT.FOOTNOTES: [], RT.ENDNOTES: []}
    for rel in document.part.rels.values():
        if rel.reltype in parts_by_type and not rel.is_external:
            parts_by_type[rel.reltype].append(rel.target_part)

    def part_root(part):
        element = getattr(part, "element", None)
        return element if element is not None else parse_xml(part.blob)

    blocks: list[str] = []
    seen: set[str] = set()

    def add_once(texts: list[str]) -> None:
        text = "\n".join(t for t in texts if t.strip())
        if text and text not in seen:  # a header repeated in every section is read once
            seen.add(text)
            blocks.append(budget.add(text))

    for part in parts_by_type[RT.HEADER]:
        add_once(_block_texts(part_root(part)))
    blocks.extend(budget.add(t) for t in _block_texts(document.element.body))
    for part in parts_by_type[RT.FOOTER]:
        add_once(_block_texts(part_root(part)))
    for reltype in (RT.FOOTNOTES, RT.ENDNOTES):
        for part in parts_by_type[reltype]:
            for note in part_root(part).iterchildren():
                # ids -1 and 0 are the separator lines Word adds itself
                if note.get(W + "id") not in ("-1", "0"):
                    add_once(_block_texts(note))
    return "\n".join(blocks)


def _block_texts(element) -> list[str]:
    """Paragraphs and tables under a body, header, footer, note or cell."""
    out: list[str] = []
    for child in element.iterchildren():
        if child.tag == W + "p":
            out.append(_paragraph_text(child))
        elif child.tag == W + "tbl":
            out.append(_table_text(child))
        elif child.tag in (W + "sdt", W + "customXml", W + "ins"):
            content = child.find(W + "sdtContent") if child.tag == W + "sdt" else child
            if content is not None:
                out.extend(_block_texts(content))
    return out


def _paragraph_text(p) -> str:
    """Text of a paragraph including runs in hyperlinks, content controls,
    tracked insertions and text boxes; w:delText (deletions) and field codes
    are different elements and are skipped."""
    parts: list[str] = []
    for el in p.iter():
        if el.tag == W + "t":
            parts.append(el.text or "")
        elif el.tag == W + "tab":
            parts.append("\t")
        elif el.tag in (W + "br", W + "cr"):
            parts.append("\n")
        elif el.tag == W + "p" and el is not p:  # paragraph inside a text box
            parts.append("\n")
    return "".join(parts)


def _children(element, tag: str):
    """Direct children with this tag, looking through content controls that
    wrap table rows or cells."""
    for child in element.iterchildren():
        if child.tag == tag:
            yield child
        elif child.tag == W + "sdt":
            content = child.find(W + "sdtContent")
            if content is not None:
                yield from _children(content, tag)


def _table_text(tbl) -> str:
    rows = []
    for tr in _children(tbl, W + "tr"):
        cells = [" ".join(t for t in _block_texts(tc) if t).replace("\n", " ") for tc in _children(tr, W + "tc")]
        rows.append("\t".join(cells))
    return "\n".join(rows)


def _xlsx(data: bytes, budget: TextBudget) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    cells = 0
    try:
        for sheet in workbook.worksheets:
            # read-only mode trusts the sheet's <dimension>, which some
            # exporters write wrongly; Excel itself ignores it
            sheet.reset_dimensions()
            parts.append(budget.add(f"# {sheet.title}"))
            for row in sheet.iter_rows(values_only=True):
                cells += max(len(row), 1)
                if cells > MAX_SPREADSHEET_CELLS:
                    raise TooLargeError(f"spreadsheet has more than {MAX_SPREADSHEET_CELLS} cells")
                values = ["" if v is None else str(v) for v in row]
                if any(values):
                    parts.append(budget.add("\t".join(values).rstrip("\t")))
    finally:
        workbook.close()
    return "\n".join(parts)
